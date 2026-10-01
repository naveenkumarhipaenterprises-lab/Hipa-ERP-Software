import csv
import io
from datetime import datetime, time

from django.conf import settings
from django.core.mail import get_connection, EmailMessage
from django.db import transaction
from django.db.models import Count, Max, Q, Sum
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import status
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response

from apps.core import periods
from apps.core.exceptions import NotConfigured
from apps.core.files import csv_response
from apps.core.metrics import choices, kpi, label_choices, num, resolve_choice
from apps.core.views import ModuleAPIView
from apps.ai_assistant import insights
from services import audit

from .models import Customer, CustomerOffer
from .serializers import CustomerWriteSerializer, customer_row

NOT_CANCELLED = ~Q(orders__status="cancelled")
TEMPLATE_HEADER = ["name", "type", "contact_person", "phone", "email", "city", "address"]


def end_of(day):
    return timezone.make_aware(datetime.combine(day, time.max))


def with_totals(qs):
    return qs.annotate(
        total_orders=Count("orders", filter=NOT_CANCELLED, distinct=True),
        total_purchase=Sum("orders__total_amount", filter=NOT_CANCELLED),
        last_order_date=Max("orders__order_date", filter=NOT_CANCELLED),
    )


def filtered_customers(params):
    qs = Customer.objects.all()
    q = (params.get("search") or "").strip()
    if q:
        qs = qs.filter(Q(name__icontains=q) | Q(city__icontains=q) | Q(contact_person__icontains=q)
                       | Q(phone__icontains=q) | Q(email__icontains=q))
    type_ = resolve_choice(Customer.Type, params.get("type"), "type")
    if type_:
        qs = qs.filter(type=type_)
    status_ = resolve_choice(Customer.Status, params.get("status"), "status")
    if status_:
        qs = qs.filter(status=status_)
    return qs


def email_channel_ready():
    return bool(settings.EMAIL_HOST) and "smtp" in settings.EMAIL_BACKEND


class CustomersModuleView(ModuleAPIView):
    module = "customers"


class OptionsView(CustomersModuleView):
    def get(self, request):
        return Response({
            "types": choices(Customer.Type),
            "statuses": label_choices(Customer.Status),
            "offer_channels": [{"value": "email", "label": "E-mail"}] if email_channel_ready() else [],
        })


class OverviewView(CustomersModuleView):
    def get(self, request):
        cur, prev = periods.resolve(self.param("range"))
        type_ = resolve_choice(Customer.Type, self.param("type"), "type")
        base = Customer.objects.filter(type=type_) if type_ else Customer.objects.all()

        def count_at(day, extra=None):
            qs = base.filter(created_at__lte=end_of(day))
            return qs.filter(**(extra or {})).count()

        by_type = []
        for value, label in Customer.Type.choices:
            if type_ and value != type_:
                continue
            by_type.append({"type": value, "label": label, **kpi(count_at(cur.end, {"type": value}), count_at(prev.end, {"type": value}))})

        in_range = Q(orders__order_date__range=(cur.start, cur.end)) & NOT_CANCELLED
        top = (base.annotate(amount=Sum("orders__total_amount", filter=in_range))
               .filter(amount__gt=0).order_by("-amount")[:5])
        return Response({
            "kpis": {"total": kpi(count_at(cur.end), count_at(prev.end)), "by_type": by_type},
            "type_distribution": [
                {"name": Customer.Type(r["type"]).label, "value": r["n"]}
                for r in base.values("type").annotate(n=Count("id")).order_by("-n")
            ],
            "locations": [
                {"city": r["city"], "customers": r["n"]}
                for r in base.values("city").annotate(n=Count("id")).order_by("-n", "city")[:10]
            ],
            "top_customers": [{"id": c.id, "name": c.name, "amount": num(c.amount)} for c in top],
            "insights": insights.texts("customers"),
        })


class GrowthView(CustomersModuleView):
    def get(self, request):
        period = periods.parse_choice(self.param("period"), ["last_12_months", "last_6_months"], "period", "last_12_months")
        type_ = resolve_choice(Customer.Type, self.param("type"), "type")
        base = Customer.objects.filter(type=type_) if type_ else Customer.objects.all()
        months = periods.last_n_months(12 if period == "last_12_months" else 6)
        if not base.exists():
            return Response([])
        return Response([{"label": label, "customers": base.filter(created_at__lte=end_of(end)).count()} for _s, end, label in months])


class CustomerListView(CustomersModuleView):
    def get(self, request):
        qs = with_totals(filtered_customers(request.query_params)).order_by("name", "id")
        return self.paginated(qs, customer_row)

    def post(self, request):
        data = request.data.copy()
        data["type"] = resolve_choice(Customer.Type, data.get("type"), "type") or ""
        data.pop("status", None)
        ser = CustomerWriteSerializer(data=data)
        ser.is_valid(raise_exception=True)
        customer = ser.save(created_by=request.user)
        audit.record(request, "Added customer", customer.name)
        return Response(customer_row(with_totals(Customer.objects.filter(pk=customer.pk)).get()), status=status.HTTP_201_CREATED)


class CustomerDetailView(CustomersModuleView):
    def get(self, request, pk):
        return Response(customer_row(get_object_or_404(with_totals(Customer.objects.all()), pk=pk)))

    def patch(self, request, pk):
        customer = get_object_or_404(Customer, pk=pk)
        data = request.data.copy()
        if "type" in data:
            data["type"] = resolve_choice(Customer.Type, data.get("type"), "type") or ""
        if "status" in data:
            data["status"] = resolve_choice(Customer.Status, data.get("status"), "status") or ""
        ser = CustomerWriteSerializer(customer, data=data, partial=True)
        ser.is_valid(raise_exception=True)
        ser.save()
        audit.record(request, "Updated customer", customer.name)
        return Response(customer_row(with_totals(Customer.objects.filter(pk=pk)).get()))


class ExportView(CustomersModuleView):
    def get(self, request):
        qs = with_totals(filtered_customers(request.query_params)).order_by("name")
        rows = (
            [c.name, c.get_type_display(), c.contact_person, c.phone, c.email, c.city, c.address,
             c.total_orders, num(c.total_purchase or 0), c.last_order_date, c.get_status_display()]
            for c in qs
        )
        header = ["Name", "Type", "Contact person", "Phone", "Email", "City", "Address",
                  "Total orders", "Total purchase (INR)", "Last order date", "Status"]
        return csv_response(f"hipa-customers-{periods.today():%Y%m%d}.csv", header, rows)


class ImportTemplateView(CustomersModuleView):
    def get(self, request):
        return csv_response("hipa-customers-template.csv", TEMPLATE_HEADER, [])


def read_rows(upload):
    """Rows of dicts from an uploaded CSV or Excel file, keyed by lower-case header."""
    name = (upload.name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        sheet = load_workbook(upload, read_only=True, data_only=True).active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        header = [str(h or "").strip().lower() for h in values[0]]
        return [{header[i]: ("" if v is None else str(v)).strip() for i, v in enumerate(r) if i < len(header)} for r in values[1:]]
    if not name.endswith(".csv"):
        raise ValidationError({"file": ["Upload a .csv or .xlsx file."]})
    try:
        text = upload.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValidationError({"file": ["The file must be UTF-8 encoded CSV."]})
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in reader]


class ImportView(CustomersModuleView):
    def post(self, request):
        upload = request.FILES.get("file")
        if not upload:
            raise ValidationError({"file": ["Choose a file to import."]})
        rows = read_rows(upload)
        if rows and not {"name", "type", "city"} <= set(rows[0].keys()):
            raise ValidationError({"file": ["The file must have the columns: " + ", ".join(TEMPLATE_HEADER)]})

        created = updated = skipped = 0
        errors = []
        for number, row in enumerate(rows, start=2):  # row 1 is the header
            if not any(row.values()):
                skipped += 1
                continue
            try:
                row_type = resolve_choice(Customer.Type, row.get("type"), "type")
            except ValidationError as exc:
                errors.append({"row": number, "message": exc.detail["type"][0]})
                continue
            data = {k: row.get(k, "") for k in TEMPLATE_HEADER}
            data["type"] = row_type or ""
            data["phone"] = data["phone"].replace(" ", "").replace("-", "")
            existing = Customer.objects.filter(name__iexact=data["name"], city__iexact=data["city"]).first()
            ser = CustomerWriteSerializer(existing, data=data, partial=bool(existing))
            if not ser.is_valid():
                field, messages = next(iter(ser.errors.items()))
                errors.append({"row": number, "message": f"{field}: {messages[0]}"})
                continue
            with transaction.atomic():
                ser.save(**({} if existing else {"created_by": request.user}))
            if existing:
                updated += 1
            else:
                created += 1
        audit.record(request, "Imported customers", f"{created} created, {updated} updated")
        return Response({"created": created, "updated": updated, "skipped": skipped, "errors": errors[:100]})


class OfferView(CustomersModuleView):
    write_module = "customer_offers"

    def post(self, request):
        segment = (request.data.get("segment") or "").strip()
        channel = (request.data.get("channel") or "").strip()
        message = (request.data.get("message") or "").strip()
        errors = {}
        if segment != "all" and segment not in Customer.Type.values:
            errors["segment"] = ["Choose a customer segment."]
        if channel != "email":
            errors["channel"] = ["Choose a connected channel."]
        if not message:
            errors["message"] = ["Write the offer message."]
        elif len(message) > 2000:
            errors["message"] = ["Keep the message under 2000 characters."]
        if errors:
            raise ValidationError(errors)
        if not email_channel_ready():
            raise NotConfigured("E-mail is not set up on the server, so offers can't be sent.")

        qs = Customer.objects.filter(status=Customer.Status.ACTIVE).exclude(email="")
        if segment != "all":
            qs = qs.filter(type=segment)
        recipients = list(qs.values_list("email", flat=True).distinct())
        if not recipients:
            raise ValidationError({"segment": ["No active customers in this segment have an e-mail address."]})

        with get_connection() as conn:
            conn.send_messages([
                EmailMessage("An offer from HIPA MASALA", message, settings.DEFAULT_FROM_EMAIL, [to], connection=conn)
                for to in recipients
            ])
        CustomerOffer.objects.create(segment=segment, channel=channel, message=message, recipients=len(recipients), sent_by=request.user)
        audit.record(request, "Sent customer offer", f"{segment} via {channel} to {len(recipients)}")
        return Response({"queued": len(recipients)})
