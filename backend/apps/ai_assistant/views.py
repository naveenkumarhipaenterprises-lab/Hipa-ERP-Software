import csv
import io

from django.db import transaction
from django.shortcuts import get_object_or_404
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response

from apps.core.exceptions import NotConfigured
from apps.core.roles import MODULE_READ, can_read, role_of
from apps.core.views import ModuleAPIView
from services import ai_client

from .context import company_summary, has_business_data
from .models import Conversation, Insight, Message

MAX_ATTACHMENT_CHARS = 20000
SYSTEM_PROMPT = (
    "You are the assistant inside the HIPA MASALA business portal (an Indian spice manufacturer). "
    "Answer clearly and briefly. When company data is provided, base every figure on it and say when "
    "the data needed to answer is missing. Never invent, estimate or assume business numbers. "
    "If the company data says there are no business records yet, tell the user plainly that there is "
    "no data yet, and do not give example figures. Amounts are in Indian Rupees (INR)."
)
NO_DATA_NOTE = ("Note: there is no business data in HIPA MASALA yet, so I can't answer from company figures. "
                "Answers will use your real records once they are entered.")

# Question starters offered on the assistant's home screen, per module the user can open
PROMPTS = {
    "sales": "How are sales this month compared with last month?",
    "inventory": "Which products are low on stock right now?",
    "finance": "Summarise this month's revenue and expenses.",
    "production": "What is in production at the moment?",
    "quality": "How many quality tests failed this month?",
    "customers": "How many active customers do we have?",
}


class AIView(ModuleAPIView):
    module = "ai_assistant"


class StatusView(AIView):
    def get(self, request):
        if not ai_client.is_configured():
            return Response({"available": False,
                             "message": "An administrator can connect an AI engine in the server settings."})
        return Response({
            "available": True,
            "models": [{"value": ai_client.model_name(), "label": ai_client.model_name()}],
            "features": {"web_search": False, "attachments": True},
        })


def visible_insights(user):
    role = role_of(user)
    modules = [m for m, roles in MODULE_READ.items() if role in roles]
    return Insight.objects.filter(module__in=modules)


class HomeView(AIView):
    def get(self, request):
        available = ai_client.is_configured()
        return Response({
            "suggestions": [q for m, q in PROMPTS.items() if can_read(request.user, m)] if available else [],
            "insights": [{"id": i.id, "kind": i.kind, "title": i.title, "text": i.text, "created_at": i.created_at}
                         for i in visible_insights(request.user)[:6]],
            "conversations": [{"id": c.id, "title": c.title, "updated_at": c.updated_at}
                              for c in Conversation.objects.filter(user=request.user)[:10]],
        })


def attachment_text(upload):
    name = (upload.name or "").lower()
    if upload.size > 5 * 1024 * 1024:
        raise ValidationError({"file": ["Attachments can be up to 5 MB."]})
    if name.endswith((".csv", ".txt")):
        try:
            return upload.read().decode("utf-8-sig")[:MAX_ATTACHMENT_CHARS]
        except UnicodeDecodeError:
            raise ValidationError({"file": ["The file must be UTF-8 text."]})
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        buf = io.StringIO()
        writer = csv.writer(buf)
        sheet = load_workbook(upload, read_only=True, data_only=True).active
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if v is None else v for v in row])
            if buf.tell() > MAX_ATTACHMENT_CHARS:
                break
        return buf.getvalue()[:MAX_ATTACHMENT_CHARS]
    raise ValidationError({"file": ["Attach a CSV, text or Excel (.xlsx) file."]})


def as_bool(value):
    return str(value).lower() in ("1", "true", "yes", "on")


class ChatView(AIView):
    def post(self, request):
        if not ai_client.is_configured():
            raise NotConfigured("The AI assistant is not connected yet.")
        message = str(request.data.get("message") or "").strip()
        if not message:
            raise ValidationError({"message": ["Type a question."]})
        if len(message) > 4000:
            raise ValidationError({"message": ["Keep the question under 4000 characters."]})
        if as_bool(request.data.get("search_web")):
            raise ValidationError({"search_web": ["Web search is not available."]})

        conversation = None
        conv_id = request.data.get("conversation_id")
        if conv_id not in (None, ""):
            conversation = get_object_or_404(Conversation, pk=conv_id, user=request.user)

        upload = request.FILES.get("file")
        attached = attachment_text(upload) if upload else ""

        history = [{"role": m.role, "content": m.text} for m in conversation.messages.all()][-20:] if conversation else []
        parts, sources = [], []
        use_data = as_bool(request.data.get("use_company_data", True))
        data_exists = use_data and has_business_data(request.user)
        if use_data and data_exists:
            parts.append(f"Company data (JSON, real records from the HIPA MASALA database):\n{company_summary(request.user)}")
            sources.append({"title": "HIPA MASALA business data"})
        elif use_data:
            parts.append("Company data: the HIPA MASALA database has no business records yet "
                         "(no sales, stock, customers, finance or other records).")
        if attached:
            parts.append(f"Attached file {upload.name}:\n{attached}")
            sources.append({"title": f"Attached file: {upload.name}"})
        parts.append(f"Question: {message}")
        reply = ai_client.complete(SYSTEM_PROMPT, history + [{"role": "user", "content": "\n\n".join(parts)}])
        if use_data and not data_exists and not attached:
            reply = f"{NO_DATA_NOTE}\n\n{reply}"

        with transaction.atomic():
            if not conversation:
                conversation = Conversation.objects.create(user=request.user, title=message[:120])
            Message.objects.create(conversation=conversation, role=Message.Role.USER, text=message,
                                   attachment_name=upload.name[:255] if upload else "")
            Message.objects.create(conversation=conversation, role=Message.Role.ASSISTANT, text=reply)
            conversation.save(update_fields=["updated_at"])
        return Response({"conversation_id": conversation.id, "reply": reply, "sources": sources})


class ConversationView(AIView):
    def get(self, request, pk):
        c = get_object_or_404(Conversation, pk=pk, user=request.user)
        return Response({"id": c.id, "title": c.title,
                         "messages": [{"role": m.role, "text": m.text, "created_at": m.created_at} for m in c.messages.all()]})
