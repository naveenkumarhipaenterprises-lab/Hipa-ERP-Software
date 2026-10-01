"""Turns a report ({ title, breakdown, table }) into CSV, Excel or PDF bytes."""
import csv
import io
from datetime import date, datetime

MIME = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _cell(value, fmt=None):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d-%m-%Y")
    return value


def to_csv(report):
    buf = io.StringIO()
    w = csv.writer(buf)
    cols = report["table"]["columns"]
    w.writerow([c["header"] for c in cols])
    for row in report["table"]["rows"]:
        w.writerow([_cell(row.get(c["key"])) for c in cols])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([report.get("title", "")])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    cols = report["table"]["columns"]
    ws.append([c["header"] for c in cols])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in report["table"]["rows"]:
        ws.append([_cell(row.get(c["key"])) for c in cols])
    for i, c in enumerate(cols, start=1):
        if c.get("format") == "inr":
            for r in ws.iter_rows(min_row=4, min_col=i, max_col=i):
                r[0].number_format = '"₹"#,##0.00'
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = max(12, min(40, len(c["header"]) + 6))
    if report.get("breakdown"):
        bs = wb.create_sheet("Breakdown")
        bs.append([report["breakdown"]["title"]])
        bs["A1"].font = Font(bold=True)
        for d in report["breakdown"]["data"]:
            bs.append([d["name"], d["value"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _latin(text):
    # Built-in PDF fonts are Latin-1 only
    return str(text).replace("₹", "Rs.").replace("—", "-").replace("–", "-").encode("latin-1", "replace").decode("latin-1")


def _pdf_value(value, fmt):
    if value is None or value == "":
        return "-"
    if isinstance(value, (date, datetime)):
        return value.strftime("%d-%m-%Y")
    if fmt == "inr":
        return f"Rs. {value:,.2f}"
    if fmt == "kg":
        return f"{value:,} kg"
    if fmt == "percent":
        return f"{value}%"
    if fmt == "number":
        return f"{value:,}"
    return str(value)


def to_pdf(report):
    from fpdf import FPDF

    cols = report["table"]["columns"]
    pdf = FPDF(orientation="L" if len(cols) > 5 else "P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 10, _latin("HIPA MASALA"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, _latin(report.get("title", "Report")), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(0, 6, _latin(f"Generated {datetime.now():%d-%m-%Y %H:%M}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    if report.get("breakdown"):
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 7, _latin(report["breakdown"]["title"]), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        for d in report["breakdown"]["data"]:
            pdf.cell(0, 5, _latin(f"{d['name']}: {_pdf_value(d['value'], report['breakdown'].get('format'))}"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

    rows = report["table"]["rows"]
    if not rows:
        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 8, "No records in this period.", new_x="LMARGIN", new_y="NEXT")
    else:
        with pdf.table(text_align="LEFT", line_height=5, first_row_as_headings=True) as table:
            pdf.set_font("Helvetica", "", 8)
            head = table.row()
            for c in cols:
                head.cell(_latin(c["header"]))
            for r in rows:
                line = table.row()
                for c in cols:
                    line.cell(_latin(_pdf_value(r.get(c["key"]), c.get("format"))))
    return bytes(pdf.output())


EXPORTERS = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}


def export(report, fmt):
    return EXPORTERS[fmt](report), MIME[fmt]
